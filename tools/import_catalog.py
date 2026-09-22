"""Regenera los dos archivos de catálogo de Código Claro. Requiere openpyxl."""
import sys, json, hashlib, datetime, pathlib, re
import openpyxl

source = pathlib.Path(sys.argv[1]) if len(sys.argv) == 2 else None
if source is None:
    raise SystemExit('Uso: python tools/import_catalog.py "catalogo.xlsx"')
root = pathlib.Path(__file__).resolve().parents[1]
w = openpyxl.load_workbook(source, read_only=True, data_only=True)
c = {'source': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'imported': datetime.date.today().isoformat(), 'version': 'Importación'}
for key, name, width, cols in [('groups','grupos',None,2),('segments','segmentos',2,3),('families','familias',4,3),('classes','clases',6,3),('products','productosSUNAT',8,3)]:
    if name not in w.sheetnames:
        raise ValueError('Falta la hoja ' + name)
    c[key] = [list(row[:cols]) for row in w[name].iter_rows(min_row=2, values_only=True) if any(v is not None for v in row)]
    ids = set()
    for row in c[key]:
        code = str(row[0])
        if code in ids or (width and not re.fullmatch(r'\d{' + str(width) + '}',code)) or not str(row[-1] or '').strip():
            raise ValueError('Código duplicado o registro inválido: ' + code)
        ids.add(code)
    if not ids:
        raise ValueError('Hoja vacía: ' + name)
for child,parent in [('segments','groups'),('families','segments'),('classes','families'),('products','classes')]:
    ids = {str(r[0]) for r in c[parent]}
    if any(str(r[1]) not in ids for r in c[child]):
        raise ValueError('Relación incompleta en ' + child)
catalog_text = 'window.CATALOG=' + json.dumps(c,ensure_ascii=False,separators=(',',':')) + ';'
codes_text = json.dumps([str(r[0]) for r in c['products']])
(root/'assets/catalog.js').write_text(catalog_text,encoding='utf-8')
(root/'server/catalog_codes.json').write_text(codes_text,encoding='utf-8')
print(f"Catálogo validado: {len(c['products'])} códigos. Archivos regenerados.")
